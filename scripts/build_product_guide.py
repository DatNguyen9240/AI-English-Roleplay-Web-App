"""Build a reader-friendly DOCX from the numbered Markdown documentation.

Design preset: compact_reference_guide
First-page pattern: editorial_cover
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "exports" / "AI-Tutor-Product-and-Architecture-Guide.docx"
SOURCE_FILES = [DOCS / f"{number:02d}-{name}.md" for number, name in [
    (1, "product-and-principles"),
    (2, "user-experience"),
    (3, "capabilities-and-content"),
    (4, "learning-and-knowledge"),
    (5, "ai-and-assessment"),
    (6, "architecture-v1"),
    (7, "contracts-data-and-events"),
    (8, "speaking-slice-current-state"),
    (9, "roadmap"),
    (10, "operations-and-quality"),
    (11, "glossary"),
]]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5B6470"


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = table_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        table_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.right_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AI Tutor · Product & Architecture Guide")
    set_font(run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(footer)
    for run in footer.runs:
        set_font(run, size=8.5, color=MUTED)


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("PRODUCT & ARCHITECTURE GUIDE")
    set_font(run, size=11, color=BLUE, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("AI Tutor")
    set_font(run, size=30, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Learning platform foundation, product flows and Architecture V1.0")
    set_font(run, size=14, color=MUTED)
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.55)
    summary.paragraph_format.right_indent = Inches(0.55)
    run = summary.add_run("A practical guide to what the product is building, how a learner moves through it, and what is implemented today.")
    set_font(run, size=11, color=INK)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Source of truth: numbered Markdown files in /docs")
    set_font(run, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_reading_map(doc):
    doc.add_heading("How to use this guide", level=1)
    doc.add_paragraph("Read the sections in order for the complete story. Section 08 is the fastest way to see what is already working in code; Section 09 is the delivery sequence.")
    rows = [
        ("01-05", "Product, learner experience, content, knowledge and AI"),
        ("06-07", "Architecture rules, contracts, data and events"),
        ("08", "Current Speaking: Car vertical slice"),
        ("09-10", "Roadmap, reliability and quality gates"),
        ("11", "Shared vocabulary"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Section"
    table.rows[0].cells[1].text = "Use"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=10, color=INK, bold=True)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                set_font(run, size=10, color=INK)
    set_table_geometry(table, [1800, 7560])
    doc.add_page_break()


def clean_inline(text):
    return re.sub(r"`([^`]+)`", r"\1", text).replace("**", "").replace("*", "")


def parse_table(lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            break
        cells = [clean_inline(cell.strip()) for cell in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_markdown(doc, content):
    lines = content.splitlines()
    index = 0
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                shade_paragraph(paragraph, LIGHT_GRAY)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = parse_table(table_lines)
            if rows:
                table = doc.add_table(rows=1, cols=len(rows[0]))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                    for col, value in enumerate(row):
                        cells[col].text = value
                        for run in cells[col].paragraphs[0].runs:
                            set_font(run, size=9.5, color=INK, bold=row_index == 0)
                        if row_index == 0:
                            set_cell_shading(cells[col], LIGHT_BLUE)
                count = len(rows[0])
                widths = [9360 // count] * count
                widths[-1] += 9360 - sum(widths)
                set_table_geometry(table, widths)
            continue
        match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if match:
            level = len(match.group(1))
            doc.add_heading(clean_inline(match.group(2)), level=level)
            index += 1
            continue
        if re.match(r"^-\s+", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(clean_inline(re.sub(r"^-\s+", "", stripped)))
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(clean_inline(re.sub(r"^\d+\.\s+", "", stripped)))
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith("#") or candidate.startswith("|") or candidate.startswith("```") or re.match(r"^(?:-|\d+\.)\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph()
        paragraph.add_run(clean_inline(" ".join(paragraph_lines)))


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def main():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "AI Tutor - Product & Architecture Guide"
    doc.core_properties.subject = "Architecture V1.0 and product direction"
    doc.core_properties.author = "AI Tutor Team"
    add_cover(doc)
    add_reading_map(doc)
    for source in SOURCE_FILES:
        add_markdown(doc, source.read_text(encoding="utf-8"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
